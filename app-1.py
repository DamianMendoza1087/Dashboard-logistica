"""App operativa de despachos, cancelaciones y devoluciones.

Ejecución:
    pip install streamlit pandas openpyxl plotly requests
    streamlit run app.py
"""

from __future__ import annotations

import io
import re
import unicodedata
from pathlib import Path

import pandas as pd
import plotly.express as px
import requests
import streamlit as st
from openpyxl import load_workbook


# -----------------------------------------------------------------------------
# CONFIGURACIÓN: complete una URL directa de descarga o cambie las rutas.
# En OneDrive, la URL debe devolver el archivo .xlsx, no una página HTML.
# También puede guardar estas claves en .streamlit/secrets.toml.
# -----------------------------------------------------------------------------
URL_ONEDRIVE_MAESTRO = ""
URL_ONEDRIVE_CONTROL = ""
RUTA_EXCEL_MAESTRO = "Cosecha Final Septiembre 2026.xlsx"
RUTA_EXCEL_CONTROL = "Cancelaciones y devoluciones.xlsx"

TABLAS_DESPACHOS = {
    "FlexTbl": ("Flex", "Flex"),
    "ColectaTbl": ("Colecta", "Colecta"),
    "CruzDelSurTbl": ("Cruz Del Sur", "Cruz del Sur"),
}
TABLA_CATEGORIAS = ("CategoriasTbl", "Categorias")


st.set_page_config(page_title="Control de Despachos", page_icon="📦", layout="wide")


def texto_limpio(valor) -> str:
    """Convierte errores y vacíos de Excel en texto vacío."""
    if valor is None or pd.isna(valor):
        return ""
    texto = str(valor).strip()
    if texto.upper() in {"#REF!", "#N/A", "#VALUE!", "NAN", "NONE", "NAT"}:
        return ""
    return re.sub(r"\s+", " ", texto)


def clave_columna(nombre) -> str:
    """Genera una clave comparable sin tildes ni signos."""
    nombre = texto_limpio(nombre).lower()
    nombre = unicodedata.normalize("NFKD", nombre).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", nombre).strip("_")


ALIAS_COLUMNAS = {
    "numero_de_venta": "numero_venta", "numero_venta": "numero_venta",
    "nro_de_venta": "numero_venta", "venta": "numero_venta",
    "pack_id": "pack_id", "packid": "pack_id",
    "sku": "sku", "codigo": "sku",
    "descripcion_del_pedido": "descripcion", "descripcion": "descripcion",
    "articulo": "descripcion", "producto": "descripcion",
    "cantidad": "cantidad", "unidades": "cantidad",
    "fecha_de_despacho": "fecha_despacho", "fecha_despacho": "fecha_despacho",
    "fecha_de_venta": "fecha_venta", "fecha_venta": "fecha_venta", "fecha": "fecha",
    "nombre_del_cliente": "cliente", "cliente": "cliente",
    "lugar_de_entrega": "lugar_entrega", "direccion_de_entrega": "direccion",
    "direccion": "direccion", "partido": "partido", "localidad": "localidad",
    "archivo": "archivo", "archivo_origen": "archivo",
    "canal": "canal", "tipo_de_envio": "canal",
    "categoria": "categoria", "categorias": "categoria",
    "factura": "factura", "nota_de_credito": "nota_credito",
    "nota_credito": "nota_credito", "nc": "nota_credito",
    "seguimiento": "seguimiento", "estado": "estado", "situacion": "situacion",
    "llego": "llego", "llego_al_deposito": "llego",
    "nota": "nota", "observacion": "nota", "observaciones": "nota",
    "canal_de_salida": "canal_salida",
}


def normalizar_columnas(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    df.columns = [ALIAS_COLUMNAS.get(clave_columna(c), clave_columna(c)) for c in df.columns]
    # Si dos encabezados terminan con el mismo nombre, conserva el primer dato disponible.
    if df.columns.duplicated().any():
        df = df.T.groupby(level=0, sort=False).first().T
    return df


def obtener_origen(url: str, ruta: str):
    """Devuelve bytes desde OneDrive o una ruta local."""
    if url.strip():
        respuesta = requests.get(url.strip(), timeout=60)
        respuesta.raise_for_status()
        if respuesta.content[:2] != b"PK":
            raise ValueError("La URL no devolvió un archivo XLSX. Use un enlace de descarga directa.")
        return io.BytesIO(respuesta.content)
    archivo = Path(ruta).expanduser()
    if not archivo.exists():
        raise FileNotFoundError(f"No se encontró: {archivo.resolve()}")
    return archivo


def leer_tabla_excel(origen, tabla: str, hoja_alternativa: str) -> pd.DataFrame:
    """Lee exactamente el rango de una tabla Excel; si falta, lee la pestaña."""
    wb = load_workbook(origen, read_only=False, data_only=True)
    try:
        for ws in wb.worksheets:
            if tabla in ws.tables:
                ref = ws.tables[tabla].ref
                filas = list(ws[ref])
                valores = [[celda.value for celda in fila] for fila in filas]
                return pd.DataFrame(valores[1:], columns=valores[0]).dropna(how="all")
        if hoja_alternativa not in wb.sheetnames:
            raise KeyError(f"No existe la tabla '{tabla}' ni la hoja '{hoja_alternativa}'.")
        ws = wb[hoja_alternativa]
        valores = list(ws.values)
        return pd.DataFrame(valores[1:], columns=valores[0]).dropna(how="all")
    finally:
        wb.close()


def leer_hoja(origen, hoja: str) -> pd.DataFrame:
    wb = load_workbook(origen, read_only=True, data_only=True)
    try:
        if hoja not in wb.sheetnames:
            return pd.DataFrame()
        valores = list(wb[hoja].values)
        return pd.DataFrame(valores[1:], columns=valores[0]).dropna(how="all")
    finally:
        wb.close()


def normalizar_canal(valor: str) -> str:
    valor = clave_columna(valor)
    if "flex" in valor:
        return "Flex"
    if "colecta" in valor:
        return "Colecta"
    if "cruz" in valor or valor == "cds":
        return "Cruz del Sur"
    return texto_limpio(valor).title() or "Sin canal"


def normalizar_despachos(df: pd.DataFrame, canal: str) -> pd.DataFrame:
    df = normalizar_columnas(df)
    df["canal"] = canal
    for col in ["numero_venta", "sku", "pack_id", "descripcion", "cliente",
                "lugar_entrega", "direccion", "partido", "localidad", "archivo"]:
        if col not in df:
            df[col] = ""
        df[col] = df[col].map(texto_limpio)
    df["cantidad"] = pd.to_numeric(df.get("cantidad", 1), errors="coerce").fillna(1)
    df["fecha_despacho"] = pd.to_datetime(df.get("fecha_despacho"), errors="coerce", dayfirst=True)
    df["id_compuesto"] = df["numero_venta"] + "|" + df["sku"]
    # Elimina repeticiones exactas, pero conserva distintos renglones de una venta multiítem.
    return df.drop_duplicates(subset=["id_compuesto", "pack_id", "fecha_despacho", "cantidad"])


def normalizar_control(df: pd.DataFrame, origen: str) -> pd.DataFrame:
    if df.empty:
        return df
    df = normalizar_columnas(df)
    df["origen_control"] = origen
    for col in ["numero_venta", "sku", "pack_id", "cliente", "descripcion", "factura",
                "nota_credito", "seguimiento", "estado", "situacion", "llego", "nota"]:
        if col not in df:
            df[col] = ""
        df[col] = df[col].map(texto_limpio)
    for col in [c for c in df.columns if c.startswith("fecha")]:
        df[col] = pd.to_datetime(df[col], errors="coerce", dayfirst=True)
    if "canal" in df:
        df["canal"] = df["canal"].map(normalizar_canal)
    df["id_compuesto"] = df["numero_venta"] + "|" + df["sku"]
    return df


def vacio_o_sin_nc(serie: pd.Series) -> pd.Series:
    limpio = serie.map(lambda x: clave_columna(x).upper())
    return limpio.isin({"", "SIN_NC", "NO", "PENDIENTE"})


def construir_alertas(cancelados: pd.DataFrame, devoluciones: pd.DataFrame) -> pd.DataFrame:
    """Reconstruye alertas desde datos crudos y explica por qué aparece cada fila."""
    partes = []
    if not cancelados.empty:
        c = cancelados.copy()
        seg = c["seguimiento"].ne("")
        falta_nc = c["factura"].ne("") & vacio_o_sin_nc(c["nota_credito"])
        c["motivo_alerta"] = ""
        c.loc[seg, "motivo_alerta"] = "Seguimiento pendiente"
        c.loc[falta_nc, "motivo_alerta"] = c.loc[falta_nc, "motivo_alerta"].map(
            lambda x: (x + " / " if x else "") + "Facturada sin NC"
        )
        partes.append(c[seg | falta_nc])
    if not devoluciones.empty:
        d = devoluciones.copy()
        seg = d["seguimiento"].ne("")
        sin_nc = vacio_o_sin_nc(d["nota_credito"])
        sin_reingreso = d["llego"].eq("")
        con_nota = d["nota"].ne("")
        regla = seg | sin_nc | sin_reingreso | con_nota
        motivos = []
        for i in d.index:
            m = []
            if seg.at[i]: m.append("Seguimiento pendiente")
            if sin_nc.at[i]: m.append("Sin NC")
            if sin_reingreso.at[i]: m.append("Sin reingreso a depósito")
            if con_nota.at[i]: m.append("Con nota")
            motivos.append(" / ".join(m))
        d["motivo_alerta"] = motivos
        partes.append(d[regla])
    return pd.concat(partes, ignore_index=True, sort=False) if partes else pd.DataFrame()


@st.cache_data(ttl=300, show_spinner="Leyendo y normalizando los Excel...")
def cargar_datos(url_maestro, ruta_maestro, url_control, ruta_control):
    maestro = obtener_origen(url_maestro, ruta_maestro)
    frames = []
    for tabla, (hoja, canal) in TABLAS_DESPACHOS.items():
        # Se vuelve a abrir la fuente para evitar problemas con streams en memoria.
        fuente = obtener_origen(url_maestro, ruta_maestro)
        frames.append(normalizar_despachos(leer_tabla_excel(fuente, tabla, hoja), canal))
    despachos = pd.concat(frames, ignore_index=True, sort=False)

    fuente_cat = obtener_origen(url_maestro, ruta_maestro)
    categorias = normalizar_columnas(leer_tabla_excel(fuente_cat, *TABLA_CATEGORIAS))
    for col in categorias.columns:
        categorias[col] = categorias[col].map(texto_limpio)

    fuente_control = obtener_origen(url_control, ruta_control)
    cancelados = normalizar_control(leer_hoja(fuente_control, "Cancelados"), "Cancelados")
    fuente_control = obtener_origen(url_control, ruta_control)
    devoluciones = normalizar_control(leer_hoja(fuente_control, "Devoluciones"), "Devoluciones")
    return despachos, categorias, cancelados, devoluciones


def enriquecer_categorias(despachos: pd.DataFrame, categorias: pd.DataFrame) -> pd.DataFrame:
    """Agrega categoría mediante SKU; usa descripción si la tabla no trae SKU."""
    resultado = despachos.copy()
    if categorias.empty or "categoria" not in categorias:
        resultado["categoria"] = "Sin categoría"
        return resultado
    llave = "sku" if "sku" in categorias else ("descripcion" if "descripcion" in categorias else None)
    if not llave:
        resultado["categoria"] = "Sin categoría"
        return resultado
    mapa = categorias[[llave, "categoria"]].copy().drop_duplicates(llave)
    mapa[llave] = mapa[llave].map(texto_limpio)
    resultado = resultado.merge(mapa, on=llave, how="left")
    resultado["categoria"] = resultado["categoria"].map(texto_limpio).replace("", "Sin categoría")
    return resultado


def tabla_visible(df: pd.DataFrame):
    preferidas = ["origen_control", "motivo_alerta", "canal", "numero_venta", "pack_id",
                  "sku", "descripcion", "cantidad", "fecha_despacho", "cliente", "estado",
                  "factura", "nota_credito", "seguimiento", "llego", "nota", "localidad",
                  "direccion", "archivo"]
    columnas = [c for c in preferidas if c in df.columns]
    st.dataframe(df[columnas], use_container_width=True, hide_index=True)


# Credenciales/rutas opcionales definidas en secrets tienen prioridad.
url_maestro = st.secrets.get("URL_ONEDRIVE_MAESTRO", URL_ONEDRIVE_MAESTRO)
url_control = st.secrets.get("URL_ONEDRIVE_CONTROL", URL_ONEDRIVE_CONTROL)
ruta_maestro = st.secrets.get("RUTA_EXCEL_MAESTRO", RUTA_EXCEL_MAESTRO)
ruta_control = st.secrets.get("RUTA_EXCEL_CONTROL", RUTA_EXCEL_CONTROL)

st.sidebar.title("📦 Navegación")
pagina = st.sidebar.radio("Ir a", [
    "Dashboard General de Despachos",
    "Carga de Trabajo y Categorías",
    "Control de Cancelaciones y Alertas",
    "Buscador de Pedidos",
])

try:
    despachos, categorias, cancelados, devoluciones = cargar_datos(
        url_maestro, ruta_maestro, url_control, ruta_control
    )
    despachos = enriquecer_categorias(despachos, categorias)
except Exception as exc:
    st.error(f"No se pudieron cargar los datos: {exc}")
    st.info("Revise las variables URL_ONEDRIVE_* o RUTA_EXCEL_* al comienzo de app.py.")
    st.stop()

canales = sorted(despachos["canal"].dropna().unique())
canales_sel = st.sidebar.multiselect("Canal", canales, default=canales)
filtrado = despachos[despachos["canal"].isin(canales_sel)].copy()
fechas_validas = filtrado["fecha_despacho"].dropna()
if not fechas_validas.empty:
    rango = st.sidebar.date_input(
        "Fecha de despacho", value=(fechas_validas.min().date(), fechas_validas.max().date())
    )
    if isinstance(rango, (tuple, list)) and len(rango) == 2:
        filtrado = filtrado[filtrado["fecha_despacho"].dt.date.between(rango[0], rango[1])]


if pagina == "Dashboard General de Despachos":
    st.title("Dashboard General de Despachos")
    k1, k2, k3, k4 = st.columns(4)
    k1.metric("Total pedidos", filtrado["numero_venta"].replace("", pd.NA).nunique())
    k2.metric("SKUs", filtrado["sku"].replace("", pd.NA).nunique())
    k3.metric("Unidades", f"{filtrado['cantidad'].sum():,.0f}".replace(",", "."))
    k4.metric("Líneas", len(filtrado))
    por_canal = filtrado.groupby("canal", as_index=False).agg(
        Pedidos=("numero_venta", "nunique"), Unidades=("cantidad", "sum")
    )
    st.plotly_chart(px.bar(por_canal, x="canal", y="Pedidos", color="canal",
                           text_auto=True, title="Pedidos por canal"), use_container_width=True)
    tabla_visible(filtrado.sort_values("fecha_despacho", ascending=False))

elif pagina == "Carga de Trabajo y Categorías":
    st.title("Carga de Trabajo y Categorías")
    c1, c2 = st.columns(2)
    por_categoria = (filtrado.groupby("categoria", as_index=False)["cantidad"].sum()
                     .sort_values("cantidad", ascending=False).head(15))
    c1.plotly_chart(px.bar(por_categoria, x="cantidad", y="categoria", orientation="h",
                           title="Categorías más despachadas", text_auto=True), use_container_width=True)
    ranking = (filtrado.groupby(["sku", "descripcion"], as_index=False)["cantidad"].sum()
               .sort_values("cantidad", ascending=False).head(20))
    etiqueta = ranking["sku"] + " · " + ranking["descripcion"].str.slice(0, 45)
    ranking = ranking.assign(producto=etiqueta)
    c2.plotly_chart(px.bar(ranking, x="cantidad", y="producto", orientation="h",
                           title="Ranking de productos", text_auto=True), use_container_width=True)
    st.subheader("Detalle del ranking")
    st.dataframe(ranking[["sku", "descripcion", "cantidad"]], use_container_width=True, hide_index=True)

elif pagina == "Control de Cancelaciones y Alertas":
    st.title("Control de Cancelaciones y Alertas")
    alertas = construir_alertas(cancelados, devoluciones)
    a1, a2, a3 = st.columns(3)
    a1.metric("Alertas activas", len(alertas))
    a2.metric("Cancelaciones", len(cancelados))
    a3.metric("Devoluciones", len(devoluciones))
    origen_sel = st.multiselect("Origen", ["Cancelados", "Devoluciones"],
                                default=["Cancelados", "Devoluciones"])
    if alertas.empty:
        st.success("No hay alertas según las reglas actuales.")
    else:
        alertas = alertas[alertas["origen_control"].isin(origen_sel)]
        motivos = alertas["motivo_alerta"].str.get_dummies(sep=" / ").sum().sort_values(ascending=False)
        st.plotly_chart(px.bar(x=motivos.index, y=motivos.values, labels={"x": "Motivo", "y": "Casos"},
                               title="Alertas por motivo", text_auto=True), use_container_width=True)
        tabla_visible(alertas)

else:
    st.title("Buscador Unificado de Pedidos")
    consulta = st.text_input("Número de Venta, SKU o Pack ID", placeholder="Escriba el dato completo o parcial")
    if consulta.strip():
        patron = re.escape(consulta.strip())
        mascara = pd.Series(False, index=despachos.index)
        for col in ["numero_venta", "sku", "pack_id"]:
            mascara |= despachos[col].str.contains(patron, case=False, na=False)
        encontrados = despachos[mascara]
        tabla_visible(encontrados)

        ids = set(encontrados["numero_venta"].dropna())
        controles = pd.concat([cancelados, devoluciones], ignore_index=True, sort=False)
        relacionados = controles[controles["numero_venta"].isin(ids)] if ids else controles.iloc[0:0]
        st.subheader("Estado en cancelaciones y devoluciones")
        if relacionados.empty:
            st.info("No se encontraron incidencias relacionadas.")
        else:
            tabla_visible(relacionados)
    else:
        st.info("Ingrese un identificador para consultar el despacho y su estado.")

st.caption("Los indicadores se recalculan desde las tablas base; las fórmulas #REF! y #N/A no se usan.")
